from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from sqlalchemy.orm import Session

from app.models.reference import CounterpartyDB, DetailsLLCDB
from app.services.reference_service import ReferenceService

CARD_TYPES = ("LLC", "IP")


class CardNotSupported(Exception):
    pass


def _val(value) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, (date, datetime)):
        return value.strftime("%d.%m.%Y")
    return str(value)


def _fio(person: dict | None) -> str | None:
    if not person:
        return None
    parts = [person.get("last_name"), person.get("name"), person.get("middle_name")]
    return " ".join(p for p in parts if p) or None


class CounterpartyCardService:
    def __init__(self, db: Session) -> None:
        self.db = db
        self.reference = ReferenceService(db)

    def build(self, counterparty_id: str) -> tuple[bytes, str] | None:
        """Возвращает (docx-байты, имя файла) или None, если контрагента нет."""
        counterparty = self.db.get(CounterpartyDB, counterparty_id)
        if not counterparty:
            return None
        if counterparty.type not in CARD_TYPES:
            raise CardNotSupported

        if counterparty.type == "LLC":
            data = self.reference.get_counterparty_llc(counterparty_id)
        else:
            data = self.reference.get_counterparty_ip(counterparty_id)
        if not data:
            return None

        accounts = self.reference.list_bank_accounts(counterparty_id)
        basic = data["basic_info"]
        details = data["details"]

        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("КАРТОЧКА ПРЕДПРИЯТИЯ")
        run.bold = True
        run.font.size = Pt(16)

        name = doc.add_paragraph()
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name.add_run(basic.get("full_name") or basic.get("short_name") or "").bold = True

        if counterparty.type == "LLC":
            rows = self._llc_rows(counterparty_id, basic, details, data)
        else:
            rows = self._ip_rows(basic, details, data)
        self._add_table(doc, "Реквизиты", rows)

        if data.get("additional_okved"):
            self._add_heading(doc, "Дополнительные виды деятельности (ОКВЭД)")
            for code in data["additional_okved"]:
                doc.add_paragraph(str(code), style="List Bullet")

        self._add_bank_accounts(doc, accounts)

        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run(f"Сформировано {datetime.now().strftime('%d.%m.%Y')}").italic = True

        buffer = BytesIO()
        doc.save(buffer)
        filename = f"Карточка предприятия — {basic.get('short_name') or counterparty_id}.docx"
        return buffer.getvalue(), filename

    def _llc_rows(self, counterparty_id: str, basic: dict, d: dict, data: dict) -> list[tuple]:
        director = data.get("director")
        basis = (
            self.db.query(DetailsLLCDB.director_basis)
            .filter(DetailsLLCDB.counterparties_id == counterparty_id)
            .scalar()
        )
        return [
            ("Полное наименование", basic.get("full_name")),
            ("Сокращённое наименование", basic.get("short_name")),
            ("ИНН", d.get("inn")),
            ("КПП", d.get("kpp")),
            ("ОГРН", d.get("ogrn")),
            ("Дата регистрации", d.get("date_register")),
            ("ОКПО", d.get("okpo")),
            ("ОКВЭД (основной)", d.get("okved")),
            ("ОКОПФ", d.get("okopf")),
            ("ОКФС", d.get("okfs")),
            ("ОКОГУ", d.get("okogu")),
            ("ОКАТО", d.get("okato")),
            ("ОКТМО", d.get("oktmo")),
            ("Система налогообложения", d.get("tax_system")),
            ("Юридический адрес", d.get("legal_address")),
            ("Фактический адрес", d.get("actual_address")),
            ("Почтовый адрес", d.get("postal_address")),
            ("Руководитель", _fio(director)),
            ("Должность руководителя", (director or {}).get("position")),
            ("Действует на основании", basis),
            ("Телефон", (director or {}).get("phone")),
            ("Электронная почта", (director or {}).get("email")),
        ]

    def _ip_rows(self, basic: dict, d: dict, data: dict) -> list[tuple]:
        owner = data.get("owner")
        return [
            ("Полное наименование", basic.get("full_name")),
            ("Сокращённое наименование", basic.get("short_name")),
            ("ИНН", d.get("inn")),
            ("ОГРНИП", d.get("ogrnip")),
            ("Дата регистрации", d.get("date_register")),
            ("ОКПО", d.get("okpo")),
            ("ОКВЭД (основной)", d.get("okved")),
            ("ОКОПФ", d.get("okopf")),
            ("ОКФС", d.get("okfs")),
            ("ОКОГУ", d.get("okogu")),
            ("ОКАТО", d.get("okato")),
            ("ОКТМО", d.get("oktmo")),
            ("Юридический адрес", d.get("legal_address")),
            ("Фактический адрес", d.get("actual_address")),
            ("Почтовый адрес", d.get("postal_address")),
            ("Индивидуальный предприниматель", _fio(owner)),
            ("Телефон", (owner or {}).get("phone")),
            ("Электронная почта", (owner or {}).get("email")),
        ]

    @staticmethod
    def _add_heading(doc, text: str) -> None:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(12)
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(12)

    def _add_table(self, doc, heading: str, rows: list[tuple]) -> None:
        self._add_heading(doc, heading)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = label
            cells[0].paragraphs[0].runs[0].bold = True
            cells[1].text = _val(value)
            cells[0].width = Pt(190)
            cells[1].width = Pt(300)

    def _add_bank_accounts(self, doc, accounts: list[dict]) -> None:
        self._add_heading(doc, "Банковские реквизиты")
        if not accounts:
            doc.add_paragraph("Не указаны")
            return
        accounts = sorted(accounts, key=lambda a: not a["is_main"])
        for account in accounts:
            title = account.get("account_name") or "Счёт"
            if account["is_main"]:
                title += " (основной)"
            if account["is_treasury"]:
                title += " (казначейский)"
            doc.add_paragraph().add_run(title).bold = True
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for label, value in (
                ("Банк", account.get("bank_name")),
                ("БИК", account.get("bik")),
                ("Корреспондентский счёт", account.get("correspondent_account")),
                ("Расчётный счёт", account.get("account_number")),
            ):
                cells = table.add_row().cells
                cells[0].text = label
                cells[0].paragraphs[0].runs[0].bold = True
                cells[1].text = _val(value)
                cells[0].width = Pt(190)
                cells[1].width = Pt(300)
